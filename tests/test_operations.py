import io
from datetime import date,timedelta
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from test_integration import app,client,csrf,login,kaka,sectest,accounts

def customer(c,brand='kaka',email='one@example.test'):
    token=csrf(c,'/'+brand+'/criar-conta')
    c.post('/'+brand+'/criar-conta',data=dict(csrf_token=token,name='Cliente',email=email,password='CustomerTest123!',confirm_password='CustomerTest123!'))
    assert c.post('/'+brand+'/login',data=dict(csrf_token=token,email=email,password='CustomerTest123!')).status_code==302
    return token

def booking(c):
    token=customer(c)
    data=dict(name='Cliente',phone='85988887777',service='Lavagem detalhada',vehicle='Carro',preferredDate=(date.today()+timedelta(days=2)).isoformat(),preferredPeriod='Tarde')
    assert c.post('/kaka/api/bookings',json=data,headers={'X-CSRF-Token':token}).status_code==200
    return token

def invoice(c):
    token=booking(c);login(c)
    assert c.post('/kaka/pedidos/1/pagamentos',data=dict(csrf_token=token,total='120',due=date.today().isoformat(),description='Sinal',reference='test-invoice')).status_code==302
    return token

def test_customer_request_chat_and_isolation(client):
    token=booking(client)
    assert client.get('/kaka/meus-pedidos/1').status_code==200
    assert client.post('/kaka/meus-pedidos/1',data=dict(csrf_token=token,action='message',message='<script>texto privado</script>')).status_code==302
    html=client.get('/kaka/meus-pedidos/1').text
    assert '&lt;script&gt;texto privado' in html
    other=app.test_client();customer(other,email='two@example.test')
    for path in ['/kaka/meus-pedidos/1','/kaka/conversas/1/mensagens','/kaka/pedidos/1/pagamentos']:
        assert other.get(path).status_code==404
    assert other.post('/kaka/meus-pedidos/1',data=dict(csrf_token=csrf(other),action='confirm',status='finalizado')).status_code==404
    assert client.get('/sectest/admin/agendamentos').status_code==302
    login(client)
    assert client.post('/kaka/admin/agendamentos/1',data=dict(csrf_token=token,action='confirm',status='agendado',total='250',confirmed_date='2026-12-20T14:00')).status_code==302
    assert 'Confirmado para' in client.get('/kaka/meus-pedidos/1').text
    assert client.post('/kaka/admin/agendamentos/1',data=dict(csrf_token=token,action='message',message='Confirmado pela equipe')).status_code==302
    assert client.get('/kaka/conversas/1/mensagens').json['messages'][-1]['administrator']==1

def test_finance_partial_duplicate_export_and_reverse(client):
    token=invoice(client)
    data=dict(csrf_token=token,action='receive',amount='20',paid_on=date.today().isoformat(),method='pix',reference='once')
    assert client.post('/kaka/admin/faturamento/1',data=data).status_code==302
    assert client.post('/kaka/admin/faturamento/1',data=data).status_code==302
    with app.app_context():assert kaka.get_db().execute('SELECT count(*) FROM receipts').fetchone()[0]==1
    for extension in ['csv','xlsx']:
        response=client.get('/kaka/admin/faturamento/exportar.'+extension+'?status=parcial')
        assert response.status_code==200
        if extension=='xlsx':
            ws=load_workbook(io.BytesIO(response.data)).active
            assert ws['E2'].value==120 and ws['F2'].value==20 and ws['G2'].value==100
    assert client.post('/kaka/admin/faturamento/1',data={**data,'reference':'bad','amount':'101'}).status_code==200
    assert client.post('/kaka/admin/faturamento/1',data=dict(csrf_token=token,action='reverse',receipt='1')).status_code==302
    assert 'R$ 120,00' in client.get('/kaka/admin/faturamento/1').text
    assert client.get('/sectest/admin/faturamento/exportar.csv').status_code==302
    assert client.get('/kaka/admin/faturamento?from=oops').status_code==400

def test_export_formula_protection(client):
    from finance import export_response
    row=dict(reference='=1+1',description=' @SUM(A1)',due='2026-01-01',status='pago',total=100,paid=100,balance=0)
    with app.test_request_context():
        response=export_response([row],'kaka','xlsx')
        ws=load_workbook(io.BytesIO(response.data)).active
        assert ws['A2'].data_type=='s' and ws['A2'].value.startswith("'")
        assert ws['B2'].value.startswith("'")

def test_provider_credit_uses_authoritative_reference_amount_and_idempotency(client):
    from client_payments import apply_payment
    invoice(client)
    with app.app_context():
        conn=kaka.get_db();conn.execute("INSERT INTO payment_attempts(reference,invoice_id,amount) VALUES ('FKunique',1,12000)");conn.commit()
        payment=dict(id=123,external_reference='FKunique',transaction_amount=120,currency_id='BRL',status='approved',date_approved='2026-10-08')
        assert not apply_payment(conn,{**payment,'transaction_amount':119})
        assert not apply_payment(conn,{**payment,'external_reference':'VORTEX123'})
        assert apply_payment(conn,payment);assert apply_payment(conn,payment)
        assert conn.execute('SELECT count(*) FROM receipts').fetchone()[0]==1
        assert apply_payment(conn,{**payment,'status':'refunded'})
        assert conn.execute('SELECT reversed FROM receipts').fetchone()[0]==1

def test_plans_estimates_monthly_idempotence_and_demo(client):
    token=customer(client,'sectest')
    data=dict(csrf_token=token,plan='pequeno',visits='100',people='2',sites='1',revenue='0',attacked='nao',sensitive='nao',action='submit')
    assert client.post('/sectest/diagnostico',data=data).status_code==302
    login(client,'sectest')
    assert client.post('/sectest/admin/contratos',data=dict(csrf_token=token,ident='1',action='activate',price='149.90')).status_code==302
    issue=dict(csrf_token=token,ident='1',action='issue',month='2026-10')
    assert client.post('/sectest/admin/contratos',data=issue).status_code==302
    assert client.post('/sectest/admin/contratos',data=issue).status_code==302
    with app.app_context():assert sectest.get_db().execute('SELECT count(*) FROM invoices').fetchone()[0]==1
    response=client.get('/sectest/pagamento/1');assert response.status_code==200
    assert 'demonstração' in response.text and 'DEMONSTRACAO-sectest-1' in response.text
    from sectest_plans import estimate
    assert estimate(dict(period='mensal',price=3990,slug='pessoal'),dict(visits=11000,people=1,revenue=0)) is None

def test_all_new_admin_pages_render(client):
    for brand in ['kaka','sectest']:
        login(client,brand)
        for path in ['agendamentos','faturamento']:
            assert client.get('/'+brand+'/admin/'+path).status_code==200
    for path in ['/sectest/admin/planos','/sectest/admin/contratos','/sectest/planos']:
        assert client.get(path).status_code==200

def test_customer_cannot_confirm_or_record_payments(client):
    token=booking(client)
    assert client.post('/kaka/meus-pedidos/1',data=dict(csrf_token=token,action='confirm',total='0',status='finalizado')).status_code==400
    assert client.post('/kaka/pedidos/1/pagamentos',data=dict(csrf_token=token,total='1')).status_code==403
    assert client.post('/kaka/meus-pedidos/1',data=dict(action='message',message='No CSRF')).status_code==400

def test_vortex_finance_own_auth_and_exports(client):
    assert client.get('/vortex7/admin/faturamento').status_code==302
    login(client,'kaka')
    assert client.get('/vortex7/admin/faturamento').status_code==302
    from datetime import datetime,timezone
    import vortex7.common as common
    with app.app_context():
        ident=common.get_db().execute('SELECT id FROM admins LIMIT 1').fetchone()[0]
    with client.session_transaction() as s:
        s['admin_id']=ident;s['admin_at']=datetime.now(timezone.utc).isoformat()
    assert client.get('/vortex7/admin/faturamento').status_code==200
    assert client.get('/vortex7/admin/faturamento/exportar.csv').status_code==200
    response=client.get('/vortex7/admin/faturamento/exportar.xlsx')
    assert load_workbook(io.BytesIO(response.data)).active['A1'].value=='Referência'

def test_webhook_signature_and_authoritative_fetch(client,monkeypatch):
    import hashlib,hmac
    import client_payments as cp
    invoice(client)
    with app.app_context():
        kaka.get_db().execute("INSERT INTO payment_attempts(reference,invoice_id,amount) VALUES ('signed',1,12000)");kaka.get_db().commit()
    monkeypatch.setattr(cp,'settings',lambda:dict(mp_access_token='test-token',test_mode='0'))
    monkeypatch.setenv('FK_MP_WEBHOOK_SECRET','test-secret')
    monkeypatch.setattr(cp.payments,'mp_get_payment',lambda ident,token:dict(id=123,external_reference='signed',transaction_amount=120,currency_id='BRL',status='approved'))
    path='/kaka/webhook/pagamentos?data.id=123'
    anonymous=app.test_client()
    assert anonymous.post(path,json={'data':{'id':'123'}}).status_code==401
    signature=hmac.new(b'test-secret',b'id:123;request-id:request123;ts:123456;',hashlib.sha256).hexdigest()
    headers={'x-signature':'ts=123456,v1='+signature,'x-request-id':'request123'}
    assert anonymous.post(path,headers=headers,json={'data':{'id':'123'}}).status_code==200
    with app.app_context():assert kaka.get_db().execute('SELECT sum(amount) FROM receipts WHERE reversed=0').fetchone()[0]==12000
