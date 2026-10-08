"""Contas a receber independentes por empresa, valores em centavos."""
import csv
import io
import secrets
from datetime import date
from decimal import Decimal, InvalidOperation
from flask import Blueprint, Response, abort, redirect, render_template, request, url_for
from accounts import current_user, require_admin


def cents(value):
    try:
        amount = Decimal(str(value).replace(',','.'))
        if not amount.is_finite() or not 0 <= amount <= 1000000:
            raise ValueError()
        return int((amount*100).quantize(Decimal('1')))
    except (InvalidOperation,ValueError):
        raise ValueError('Informe um valor entre zero e R$ 1.000.000.')


def money(value):
    return ('R$ '+format(value/100,',.2f')).replace(',','X').replace('.',',').replace('X','.')


def init_finance(conn):
    conn.executescript('''
    CREATE TABLE IF NOT EXISTS invoices (
      id INTEGER PRIMARY KEY, request_id INTEGER NOT NULL, description TEXT NOT NULL,
      total INTEGER NOT NULL CHECK(total>0), due TEXT NOT NULL, reference TEXT UNIQUE NOT NULL,
      canceled INTEGER NOT NULL DEFAULT 0, created_at TEXT NOT NULL DEFAULT (datetime('now')));
    CREATE TABLE IF NOT EXISTS receipts (
      id INTEGER PRIMARY KEY, invoice_id INTEGER NOT NULL, amount INTEGER NOT NULL CHECK(amount>0),
      method TEXT NOT NULL, paid_on TEXT NOT NULL, actor TEXT NOT NULL,
      reference TEXT UNIQUE NOT NULL, reversed INTEGER NOT NULL DEFAULT 0,
      created_at TEXT NOT NULL DEFAULT (datetime('now')));
    CREATE INDEX IF NOT EXISTS receipt_invoice ON receipts(invoice_id);
    ''')


def invoice_rows(conn):
    rows = conn.execute('''SELECT i.*,r.owner_id,r.service,r.source_id,
      COALESCE((SELECT sum(amount) FROM receipts WHERE invoice_id=i.id AND reversed=0),0) paid
      FROM invoices i JOIN client_requests r ON r.id=i.request_id ORDER BY i.id DESC''').fetchall()
    result=[]
    for row in rows:
        d=dict(row);d['balance']=max(0,d['total']-d['paid'])
        d['status']='cancelado' if d['canceled'] else 'pago' if not d['balance'] else 'atrasado' if d['due']<date.today().isoformat() else 'parcial' if d['paid'] else 'pendente'
        result.append(d)
    return result


def filtered(rows):
    start=request.args.get('from','');end=request.args.get('to','');state=request.args.get('status','');search=request.args.get('q','').lower().strip()
    for value in (start,end):
        if value:
            try: date.fromisoformat(value)
            except ValueError: abort(400,'Data inválida.')
    if start and end and start>end: abort(400,'O início deve ser anterior ao fim.')
    return [r for r in rows if (not start or r['due']>=start) and (not end or r['due']<=end) and (not state or r['status']==state) and (not search or search in (str(r['reference'])+' '+r['description']).lower())]


def export_response(rows, brand, extension):
    headers=['Referência','Descrição','Vencimento','Status','Valor (R$)','Recebido (R$)','Saldo (R$)']
    data=[[r['reference'],r['description'],r['due'],r['status'],r['total']/100,r['paid']/100,r['balance']/100] for r in rows]
    def safe(v):
        if isinstance(v,str) and v.lstrip().startswith(('=','+','-','@')): return "'"+v
        return v
    if extension=='csv':
        out=io.StringIO();writer=csv.writer(out,delimiter=';');writer.writerow(headers)
        for row in data: writer.writerow([format(v,'.2f').replace('.',',') if isinstance(v,float) else safe(v) for v in row])
        blob=out.getvalue().encode('utf-8-sig');mime='text/csv; charset=utf-8'
    elif extension=='xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
        book=Workbook();sheet=book.active;sheet.title='Faturamento';sheet.append(headers)
        for row in data: sheet.append([safe(v) for v in row])
        for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='182431')
        for row in sheet.iter_rows(min_row=2,min_col=5,max_col=7):
            for cell in row:cell.number_format='"R$" #,##0.00'
        for i,width in enumerate([34,48,16,18,20,20,20],1):sheet.column_dimensions[get_column_letter(i)].width=width
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        out=io.BytesIO();book.save(out);blob=out.getvalue();mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    else:abort(404)
    return Response(blob,mimetype=mime,headers={'Content-Disposition':f'attachment; filename="{brand}-faturamento.{extension}"','Cache-Control':'no-store'})


def stats(rows):
    active=[r for r in rows if not r.get('canceled')]
    months={}
    for row in active:
        key=row['due'][:7];months[key]=months.get(key,0)+row['paid']
    top=max(months.values(),default=1) or 1
    return dict(total=sum(r['total'] for r in active),paid=sum(r['paid'] for r in active),balance=sum(r['balance'] for r in active),
                charts=[(key,value,round(value/top*100)) for key,value in sorted(months.items())])


def register_finance(app,brand,prefix,get_db):
    bp=Blueprint(brand+'_finance',__name__)
    def render(template,**kwargs):
        return render_template('finance/'+template,brand=brand,brand_name='FK Káka Autodetail' if brand=='kaka' else 'SecTest',admin=True,money=money,**kwargs)
    @bp.after_request
    def private(response):
        response.headers['Cache-Control']='no-store';return response

    @bp.get('/admin/faturamento')
    @require_admin(brand)
    def dashboard():
        rows=filtered(invoice_rows(get_db()))
        return render('dashboard.html',title='Faturamento',rows=rows,summary=stats(rows),vortex=False)

    @bp.get('/admin/faturamento/exportar.<extension>')
    @require_admin(brand)
    def export(extension):
        return export_response(filtered(invoice_rows(get_db())),brand,extension)

    @bp.route('/pedidos/<int:ident>/pagamentos',methods=['GET','POST'])
    def request_invoices(ident):
        from accounts import verify_csrf
        admin=current_user(brand,True);customer=current_user(brand)
        if not admin and not customer:return redirect(url_for(brand+'_accounts.login'))
        conn=get_db();item=conn.execute('SELECT * FROM client_requests WHERE id=?',(ident,)).fetchone()
        if not item or (not admin and item['owner_id']!=customer['id']):abort(404)
        error=None
        if request.method=='POST':
            verify_csrf()
            if not admin:abort(403)
            try:
                total=cents(request.form.get('total',''))
                if total<=0:raise ValueError('O valor deve ser maior que zero.')
                due=date.fromisoformat(request.form.get('due','')).isoformat()
                description=request.form.get('description','').strip()[:250]
                if not description:raise ValueError('Informe a descrição da cobrança.')
                reference=request.form.get('reference','')
                if not reference or len(reference)>80:raise ValueError('Atualize a página.')
                conn.execute('INSERT OR IGNORE INTO invoices(request_id,description,total,due,reference) VALUES (?,?,?,?,?)',
                  (ident,description,total,due,reference));conn.commit()
                return redirect(request.path)
            except ValueError as exc:error=str(exc)
        rows=[r for r in invoice_rows(conn) if r['request_id']==ident]
        return render_template('finance/invoices.html',brand=brand,brand_name='FK Káka Autodetail' if brand=='kaka' else 'SecTest',admin=bool(admin),
          title='Pagamentos do pedido #'+str(ident),rows=rows,item=item,money=money,error=error,nonce=secrets.token_hex(16),today=date.today().isoformat())

    @bp.route('/admin/faturamento/<int:ident>',methods=['GET','POST'])
    @require_admin(brand)
    def invoice(ident):
        conn=get_db();row=next((r for r in invoice_rows(conn) if r['id']==ident),None)
        if not row:abort(404)
        error=None
        if request.method=='POST':
            try:
                action=request.form.get('action')
                if action=='receive':
                    conn.execute('BEGIN IMMEDIATE')
                    row=next(r for r in invoice_rows(conn) if r['id']==ident)
                    amount=cents(request.form.get('amount',''));paid_on=date.fromisoformat(request.form.get('paid_on','')).isoformat()
                    method=request.form.get('method','');ref=request.form.get('reference','')
                    if row['canceled'] or amount<=0 or amount>row['balance']:raise ValueError('O recebimento deve ser positivo e não pode superar o saldo.')
                    if paid_on>date.today().isoformat():raise ValueError('Recebimentos não podem ter data futura.')
                    if method not in ('pix','dinheiro','cartao','boleto','demonstracao') or not ref or len(ref)>80:raise ValueError('Dados inválidos.')
                    if brand=='sectest':method='demonstracao'
                    conn.execute('INSERT OR IGNORE INTO receipts(invoice_id,amount,method,paid_on,actor,reference) VALUES (?,?,?,?,?,?)',
                      (ident,amount,method,paid_on,current_user(brand,True)['email'],ref))
                elif action=='reverse':
                    conn.execute('UPDATE receipts SET reversed=1 WHERE id=? AND invoice_id=? AND method<>?',(request.form.get('receipt'),ident,'mercadopago'))
                elif action=='cancel':
                    if row['paid']:raise ValueError('Estorne os recebimentos antes de cancelar.')
                    conn.execute('UPDATE invoices SET canceled=1 WHERE id=?',(ident,))
                else:abort(400)
                conn.commit();return redirect(request.path)
            except ValueError as exc:conn.rollback();error=str(exc)
        receipts=conn.execute('SELECT * FROM receipts WHERE invoice_id=? ORDER BY id DESC',(ident,)).fetchall()
        return render('invoice.html',title='Cobrança #'+str(ident),row=row,receipts=receipts,error=error,nonce=secrets.token_hex(16),today=date.today().isoformat())

    app.register_blueprint(bp,url_prefix=prefix)
